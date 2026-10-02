"""Hidden-oracle bounded cross-format structural-preservation preflight.

This suite generates native DOCX/XLSX/PPTX/PDF artifacts, exposes only the artifact
bytes and explicit structural/content edit request, then independently verifies that
the requested target changed while unrelated structure and format invariants remain.

It deliberately does NOT claim pixel/render equivalence or arbitrary semantic edit
target selection. Those remain terminal-scope residuals.
"""
from __future__ import annotations

import base64
from copy import deepcopy
from io import BytesIO
import random
from typing import Any, Mapping
import xml.etree.ElementTree as ET
import zipfile

FORMATS=("docx","xlsx","pptx","pdf")

def _b64(data: bytes) -> str:
    return base64.b64encode(data).decode("ascii")

def _unb64(data: str) -> bytes:
    return base64.b64decode(data.encode("ascii"),validate=True)

def _zip_members(data: bytes) -> dict[str,bytes]:
    with zipfile.ZipFile(BytesIO(data),"r") as z:
        names=[i.filename for i in z.infolist()]
        if len(names)!=len(set(names)):
            raise ValueError("DUPLICATE_MEMBER")
        return {n:z.read(n) for n in names}

def _node_at(root: ET.Element,path: list[int]) -> ET.Element:
    n=root
    for i in path:
        n=n[i]
    return n

def _find_text_path(root: ET.Element,target: str) -> list[int]:
    hits=[]
    def walk(node,path):
        if node.text==target:
            hits.append(list(path))
        for i,ch in enumerate(node):
            walk(ch,path+[i])
    walk(root,[])
    if len(hits)!=1:
        raise ValueError(f"TARGET_TEXT_NOT_UNIQUE:{len(hits)}")
    return hits[0]

def _locate_ooxml_text(data: bytes,target: str) -> tuple[str,list[int],str]:
    members=_zip_members(data)
    hits=[]
    for name,payload in members.items():
        if not (name.endswith(".xml") or name.endswith(".rels")):
            continue
        try:
            root=ET.fromstring(payload)
        except Exception:
            continue
        try:
            path=_find_text_path(root,target)
        except ValueError:
            continue
        node=_node_at(root,path)
        hits.append((name,path,node.tag))
    if len(hits)!=1:
        raise ValueError(f"PACKAGE_TARGET_NOT_UNIQUE:{len(hits)}")
    return hits[0]

def _make_docx(seed:int,target:str)->bytes:
    from docx import Document
    d=Document()
    p=d.add_paragraph()
    r=p.add_run(target); r.bold=True
    k=p.add_run(" KEEP"); k.italic=True
    p2=d.add_paragraph("UNCHANGED")
    p2.style="Quote"
    d.sections[0].top_margin=d.sections[0].top_margin
    b=BytesIO(); d.save(b); return b.getvalue()

def _make_xlsx(seed:int,target:str)->bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill
    w=Workbook(); s=w.active; s.title="Main"
    s["A1"]=target
    s["B1"]="KEEP"; s["B1"].font=Font(bold=True); s["B1"].fill=PatternFill("solid",fgColor="FFFF00")
    s["C2"]="=1+2"; s.column_dimensions["B"].width=24.0; s.row_dimensions[1].height=28.0
    s.merge_cells("D4:E4"); s["D4"]="MERGED"
    a=w.create_sheet("Archive"); a["C3"]="UNCHANGED"
    b=BytesIO(); w.save(b); return b.getvalue()

def _make_pptx(seed:int,target:str)->bytes:
    from pptx import Presentation
    from pptx.util import Inches,Pt
    p=Presentation(); slide=p.slides.add_slide(p.slide_layouts[6])
    s1=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1))
    r=s1.text_frame.paragraphs[0].add_run(); r.text=target; r.font.bold=True; r.font.size=Pt(22)
    s2=slide.shapes.add_textbox(Inches(1),Inches(3),Inches(4),Inches(1))
    r2=s2.text_frame.paragraphs[0].add_run(); r2.text="UNCHANGED"; r2.font.italic=True; r2.font.size=Pt(16)
    b=BytesIO(); p.save(b); return b.getvalue()

def _make_pdf(seed:int,target:str)->bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject,DictionaryObject,NameObject
    w=PdfWriter(); page=w.add_blank_page(width=240,height=180)
    font=DictionaryObject({
        NameObject("/Type"):NameObject("/Font"),
        NameObject("/Subtype"):NameObject("/Type1"),
        NameObject("/BaseFont"):NameObject("/Helvetica"),
    })
    font_ref=w._add_object(font)
    page[NameObject("/Resources")]=DictionaryObject({
        NameObject("/Font"):DictionaryObject({NameObject("/F1"):font_ref})
    })
    stream=DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 100 Td ({target}) Tj 0 -25 Td (UNCHANGED) Tj ET".encode("ascii"))
    page[NameObject("/Contents")]=w._add_object(stream)
    b=BytesIO(); w.write(b); return b.getvalue()

def _docx_snapshot(data:bytes)->dict:
    from docx import Document
    d=Document(BytesIO(data))
    return {
        "paragraphs":[[
            {"text":r.text,"bold":r.bold,"italic":r.italic,"style":r.style.name if r.style else None}
            for r in p.runs
        ] for p in d.paragraphs],
        "paragraph_styles":[p.style.name if p.style else None for p in d.paragraphs],
        "section_count":len(d.sections),
    }

def _xlsx_snapshot(data:bytes)->dict:
    from openpyxl import load_workbook
    w=load_workbook(BytesIO(data),data_only=False)
    sheets={}
    for s in w.worksheets:
        cells={}
        for row in s.iter_rows():
            for c in row:
                if c.value is not None:
                    cells[c.coordinate]=(c.value,c.style_id,c.number_format)
        sheets[s.title]={
            "cells":cells,
            "merged":sorted(str(x) for x in s.merged_cells.ranges),
            "width_B":s.column_dimensions["B"].width,
            "height_1":s.row_dimensions[1].height,
        }
    return {"sheetnames":w.sheetnames,"sheets":sheets}

def _pptx_snapshot(data:bytes)->dict:
    from pptx import Presentation
    p=Presentation(BytesIO(data))
    slides=[]
    for sl in p.slides:
        shapes=[]
        for sh in sl.shapes:
            runs=[]
            if getattr(sh,"has_text_frame",False):
                for par in sh.text_frame.paragraphs:
                    for r in par.runs:
                        runs.append((r.text,r.font.bold,r.font.italic,r.font.size.pt if r.font.size else None))
            shapes.append({
                "shape_type":int(sh.shape_type),
                "left":int(sh.left),"top":int(sh.top),"width":int(sh.width),"height":int(sh.height),
                "runs":runs,
            })
        slides.append(shapes)
    return {"slide_count":len(p.slides),"slides":slides}

def _pdf_snapshot(data:bytes)->dict:
    from pypdf import PdfReader
    from pypdf.generic import ContentStream
    r=PdfReader(BytesIO(data),strict=True)
    pages=[]
    for p in r.pages:
        cs=ContentStream(p.get_contents(),r)
        pages.append({
            "text":p.extract_text(),
            "mediabox":tuple(float(x) for x in p.mediabox),
            "rotation":int(p.get("/Rotate",0) or 0),
            "operators":[op.decode("latin1") for _,op in cs.operations],
        })
    return {"page_count":len(r.pages),"pages":pages}

def _snapshot(fmt:str,data:bytes)->dict:
    return {"docx":_docx_snapshot,"xlsx":_xlsx_snapshot,"pptx":_pptx_snapshot,"pdf":_pdf_snapshot}[fmt](data)

def _expected_snapshot(fmt:str,source:bytes,old:str,new:str)->dict:
    # Independent semantic expectation built through format APIs, not the Brain transaction executor.
    if fmt=="docx":
        from docx import Document
        d=Document(BytesIO(source))
        hits=[r for p in d.paragraphs for r in p.runs if r.text==old]
        if len(hits)!=1: raise ValueError("DOCX_TARGET")
        hits[0].text=new; b=BytesIO(); d.save(b); return _docx_snapshot(b.getvalue())
    if fmt=="xlsx":
        from openpyxl import load_workbook
        w=load_workbook(BytesIO(source),data_only=False)
        hits=[c for s in w.worksheets for row in s.iter_rows() for c in row if c.value==old]
        if len(hits)!=1: raise ValueError("XLSX_TARGET")
        hits[0].value=new; b=BytesIO(); w.save(b); return _xlsx_snapshot(b.getvalue())
    if fmt=="pptx":
        from pptx import Presentation
        p=Presentation(BytesIO(source)); hits=[]
        for sl in p.slides:
            for sh in sl.shapes:
                if getattr(sh,"has_text_frame",False):
                    for par in sh.text_frame.paragraphs:
                        hits.extend(r for r in par.runs if r.text==old)
        if len(hits)!=1: raise ValueError("PPTX_TARGET")
        hits[0].text=new; b=BytesIO(); p.save(b); return _pptx_snapshot(b.getvalue())
    raise ValueError("EXPECTED_ONLY_OOXML")

def generate_case(fmt:str,seed:int)->dict[str,Any]:
    if fmt not in FORMATS or not isinstance(seed,int) or isinstance(seed,bool):
        raise ValueError("INPUT")
    r=random.Random(seed)
    old=f"OLD{seed:06d}"
    new=f"NEW{seed:06d}"
    maker={"docx":_make_docx,"xlsx":_make_xlsx,"pptx":_make_pptx,"pdf":_make_pdf}[fmt]
    source=maker(seed,old)
    if fmt=="pdf":
        edit={"kind":"PDF_REPLACE_UNIQUE_TEXT","page_index":0,"old":old,"new":new}
        oracle={"old":old,"new":new,"source_snapshot":_pdf_snapshot(source)}
    else:
        member,path,tag=_locate_ooxml_text(source,old)
        edit={"kind":"XML_SET_TEXT","member":member,"path":path,"new_text":new,"expected_tag":tag}
        oracle={
            "old":old,"new":new,"target_member":member,"target_path":path,
            "source_snapshot":_snapshot(fmt,source),
            "expected_snapshot":_expected_snapshot(fmt,source,old,new),
        }
    return {
        "schema":"PROJECT_BRAIN_NATIVE_ARTIFACT_CROSS_FORMAT_INFORMATION_SAFE_PREFLIGHT_V1",
        "behavior_id":"NATIVE_ARTIFACT_STRUCTURED_EDIT_PRESERVATION_001",
        "seed":seed,
        "task":{
            "format":fmt,
            "document_b64":_b64(source),
            "edit":edit,
            "request":"Apply exactly the declared native structural/content edit and preserve unrelated structure.",
        },
        "_oracle":oracle,
    }

def public_task(case:Mapping[str,Any])->dict[str,Any]:
    return {k:v for k,v in case.items() if k!="_oracle"}

def _semantic_target_only_diff(source_xml:bytes,out_xml:bytes,path:list[int],old:str,new:str)->bool:
    try:
        a=ET.fromstring(source_xml); b=ET.fromstring(out_xml)
        an=_node_at(a,path); bn=_node_at(b,path)
        if an.text!=old or bn.text!=new or an.tag!=bn.tag:
            return False
        bn.text=old
        return ET.tostring(a,encoding="utf-8")==ET.tostring(b,encoding="utf-8")
    except Exception:
        return False

def score_case(case:Mapping[str,Any],candidate:Mapping[str,Any])->dict[str,Any]:
    if not isinstance(candidate,Mapping) or candidate.get("status")!="OK":
        return {"pass":False,"reason":"CANDIDATE_STATUS"}
    try:
        out=_unb64(candidate["output_b64"])
    except Exception:
        return {"pass":False,"reason":"OUTPUT_DECODE"}
    fmt=case["task"]["format"]; oracle=case["_oracle"]; source=_unb64(case["task"]["document_b64"])
    try:
        if fmt in {"docx","xlsx","pptx"}:
            before=_zip_members(source); after=_zip_members(out)
            if set(before)!=set(after):
                return {"pass":False,"reason":"PACKAGE_MEMBER_SET_CHANGED"}
            target=oracle["target_member"]
            for name in before:
                if name!=target and before[name]!=after[name]:
                    return {"pass":False,"reason":"UNRELATED_PACKAGE_MEMBER_MUTATED","member":name}
            if not _semantic_target_only_diff(before[target],after[target],oracle["target_path"],oracle["old"],oracle["new"]):
                return {"pass":False,"reason":"TARGET_XML_HAS_UNDECLARED_CHANGE"}
            got=_snapshot(fmt,out)
            if got!=oracle["expected_snapshot"]:
                return {"pass":False,"reason":"FORMAT_REOPEN_SNAPSHOT_MISMATCH"}
        else:
            got=_pdf_snapshot(out); before=oracle["source_snapshot"]
            if got["page_count"]!=before["page_count"]:
                return {"pass":False,"reason":"PDF_PAGE_COUNT_CHANGED"}
            bp=before["pages"][0]; gp=got["pages"][0]
            if bp["mediabox"]!=gp["mediabox"] or bp["rotation"]!=gp["rotation"] or bp["operators"]!=gp["operators"]:
                return {"pass":False,"reason":"PDF_GEOMETRY_OR_OPERATOR_SEQUENCE_CHANGED"}
            if oracle["new"] not in gp["text"] or oracle["old"] in gp["text"] or "UNCHANGED" not in gp["text"]:
                return {"pass":False,"reason":"PDF_TEXT_PRESERVATION_MISMATCH"}
    except Exception as exc:
        return {"pass":False,"reason":"ORACLE_EXCEPTION","error":type(exc).__name__+":"+str(exc)}
    return {"pass":True,"reason":"PASS"}

